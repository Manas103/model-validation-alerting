"""Tests for the independent price-verification extension.

`mvguard/marks_producer.py` builds the synthetic submitted-mark and
independent-source curves and plants 24 known violations;
`mvguard/price_verification.py` is the engine under test;
`mvguard/exception_pack.py` is the Excel writer under test.
"""

import os
import tempfile

import openpyxl
import pytest

from mvguard.exception_pack import write_exception_pack
from mvguard.marks_producer import (
    SEED_SPECS,
    generate_clean_marks,
    generate_seeded_marks,
    seed_target_snapshot_ids,
)
from mvguard.price_verification import (
    PriceVerificationRuleError,
    evaluate_stream,
    load_rules,
)


def test_rule_file_loads_all_three_families():
    rules = load_rules()
    assert set(rules) == {"staleness", "off_market", "calendar_spread"}
    assert rules["staleness"]["stale_sessions"] > 0
    assert rules["off_market"]["tolerance"] > 0
    assert rules["calendar_spread"]["tolerance"] > 0
    for rule in rules.values():
        assert rule["severity"] in ("info", "warning", "critical")


def test_rule_file_rejects_missing_required_field(tmp_path):
    bad = tmp_path / "bad.yaml"
    bad.write_text("- name: x\n  family: off_market\n", encoding="utf-8")
    with pytest.raises(PriceVerificationRuleError):
        load_rules(str(bad))


def test_rule_file_rejects_unknown_family(tmp_path):
    bad = tmp_path / "bad.yaml"
    bad.write_text("- name: x\n  family: bogus\n  tolerance: 0.01\n", encoding="utf-8")
    with pytest.raises(PriceVerificationRuleError):
        load_rules(str(bad))


def test_clean_baseline_has_no_violations():
    rules = load_rules()
    clean = generate_clean_marks()
    assert sum(len(s["tenors"]) for s in clean) == 12000
    alerts = evaluate_stream(clean, rules)
    assert alerts == []


def test_seeded_stream_catches_all_24_with_zero_false_positives():
    rules = load_rules()
    seeded = generate_seeded_marks()
    alerts = evaluate_stream(seeded, rules)

    by_snapshot = {}
    for alert in alerts:
        by_snapshot.setdefault(alert["snapshot_id"], []).append(alert)

    seeded_ids = set()
    for spec in SEED_SPECS:
        seeded_ids |= seed_target_snapshot_ids(spec)

    false_positive_ids = set(by_snapshot) - seeded_ids
    assert false_positive_ids == set(), (
        "alert(s) fired in a snapshot with no planted seed: {}".format(false_positive_ids)
    )

    missed = []
    for spec in SEED_SPECS:
        target_ids = seed_target_snapshot_ids(spec)
        families_fired = {a["family"] for sid in target_ids for a in by_snapshot.get(sid, [])}
        if spec["family"] not in families_fired:
            missed.append(spec["id"])
    assert missed == [], "seed(s) whose target family never fired: {}".format(missed)


def test_every_alert_carries_triggering_inputs():
    rules = load_rules()
    seeded = generate_seeded_marks()
    alerts = evaluate_stream(seeded, rules)
    assert len(alerts) > 0
    for alert in alerts:
        assert "snapshot" in alert and alert["snapshot"], "alert missing triggering snapshot data"
        assert "observed" in alert and alert["observed"], "alert missing observed values"
        assert "commodity" in alert["snapshot"] and "session_index" in alert["snapshot"]


def test_staleness_fires_once_per_stale_run():
    rules = load_rules()
    seeded = generate_seeded_marks()
    alerts = evaluate_stream(seeded, rules)
    stale_alerts = [a for a in alerts if a["family"] == "staleness"]
    keys = [(a["snapshot"]["commodity"], a["snapshot"]["tenor"], a["session_index"]) for a in stale_alerts]
    assert len(keys) == len(set(keys)), "staleness fired more than once for the same run"
    assert len(stale_alerts) == 8


def test_exception_pack_writes_expected_rows():
    rules = load_rules()
    seeded = generate_seeded_marks()
    alerts = evaluate_stream(seeded, rules)

    with tempfile.TemporaryDirectory() as tmp:
        path = os.path.join(tmp, "exceptions.xlsx")
        written_path = write_exception_pack(alerts, path)
        assert written_path == path
        assert os.path.exists(path)

        wb = openpyxl.load_workbook(path)
        assert set(wb.sheetnames) == {"All exceptions", "staleness", "off_market", "calendar_spread"}

        off_market_alerts = [a for a in alerts if a["family"] == "off_market"]
        assert off_market_alerts, "expected at least one off-market alert to check against"
        first = off_market_alerts[0]

        ws = wb["off_market"]
        header = [c.value for c in ws[1]]
        row = [c.value for c in ws[2]]
        row_dict = dict(zip(header, row))
        assert row_dict["rule_name"] == first["rule_name"]
        assert row_dict["commodity"] == first["snapshot"]["commodity"]
        assert row_dict["tenor"] == first["snapshot"]["tenor"]
        assert row_dict["submitted_mark"] == first["snapshot"]["submitted_mark"]
        assert row_dict["independent_mark"] == first["snapshot"]["independent_mark"]

        summary = wb["All exceptions"]
        assert summary["B2"].value == len(alerts)
