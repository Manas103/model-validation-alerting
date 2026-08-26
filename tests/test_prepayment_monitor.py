"""Tests for the prepayment-model-monitoring extension.

`mvguard/prepayment_monitor_producer.py` builds the synthetic cohort panel
and plants 30 known violations; `mvguard/prepayment_monitor.py` is the
engine under test; `mvguard/prepayment_exception_pack.py` is the Excel
writer under test.
"""

import os
import tempfile

import openpyxl
import pytest

from mvguard.prepayment_exception_pack import write_exception_pack
from mvguard.prepayment_monitor_producer import (
    COHORT_COUNT,
    MONTH_COUNT,
    SEED_SPECS,
    generate_clean_snapshots,
    generate_seeded_snapshots,
    seed_target_snapshot_ids,
)
from mvguard.prepayment_monitor import (
    PrepaymentMonitorRuleError,
    evaluate_stream,
    load_rules,
    psi,
)


def test_rule_file_loads_all_three_families():
    rules = load_rules()
    assert set(rules) == {"cpr_tolerance", "stale_input", "population_stability"}
    assert rules["cpr_tolerance"]["tolerance"] > 0
    assert rules["stale_input"]["stale_months"] > 0
    assert rules["population_stability"]["threshold"] > 0
    for rule in rules.values():
        assert rule["severity"] in ("info", "warning", "critical")


def test_rule_file_rejects_missing_required_field(tmp_path):
    bad = tmp_path / "bad.yaml"
    bad.write_text("- name: x\n  family: cpr_tolerance\n", encoding="utf-8")
    with pytest.raises(PrepaymentMonitorRuleError):
        load_rules(str(bad))


def test_rule_file_rejects_unknown_family(tmp_path):
    bad = tmp_path / "bad.yaml"
    bad.write_text("- name: x\n  family: bogus\n  tolerance: 1\n", encoding="utf-8")
    with pytest.raises(PrepaymentMonitorRuleError):
        load_rules(str(bad))


def test_panel_shape_is_15000_cohort_months():
    clean = generate_clean_snapshots()
    assert len(clean) == COHORT_COUNT * MONTH_COUNT == 15000


def test_clean_baseline_has_no_violations():
    rules = load_rules()
    clean = generate_clean_snapshots()
    alerts = evaluate_stream(clean, rules, population_size=COHORT_COUNT)
    assert alerts == []


def test_seeded_stream_catches_all_30_with_zero_false_positives():
    rules = load_rules()
    seeded = generate_seeded_snapshots()
    alerts = evaluate_stream(seeded, rules, population_size=COHORT_COUNT)

    by_snapshot = {}
    for alert in alerts:
        by_snapshot.setdefault(alert["snapshot_id"], []).append(alert)

    seeded_ids = set()
    for spec in SEED_SPECS:
        seeded_ids |= seed_target_snapshot_ids(spec)

    false_positive_ids = set(by_snapshot) - seeded_ids
    assert false_positive_ids == set(), (
        "alert(s) fired in a snapshot with no planted seed: {}".format(false_positive_ids)
    )

    missed = []
    for spec in SEED_SPECS:
        target_ids = seed_target_snapshot_ids(spec)
        families_fired = {a["family"] for sid in target_ids for a in by_snapshot.get(sid, [])}
        if spec["family"] not in families_fired:
            missed.append(spec["id"])
    assert missed == [], "seed(s) whose target family never fired: {}".format(missed)


def test_every_alert_carries_triggering_inputs():
    rules = load_rules()
    seeded = generate_seeded_snapshots()
    alerts = evaluate_stream(seeded, rules, population_size=COHORT_COUNT)
    assert len(alerts) > 0
    for alert in alerts:
        assert "snapshot" in alert and alert["snapshot"], "alert missing triggering snapshot data"
        assert "observed" in alert and alert["observed"], "alert missing observed values"


def test_stale_input_fires_once_per_stale_run():
    rules = load_rules()
    seeded = generate_seeded_snapshots()
    alerts = evaluate_stream(seeded, rules, population_size=COHORT_COUNT)
    stale_alerts = [a for a in alerts if a["family"] == "stale_input"]
    keys = [(a["cohort_id"], a["month_index"]) for a in stale_alerts]
    assert len(keys) == len(set(keys)), "stale_input fired more than once for the same run"
    assert len(stale_alerts) == 10


def test_psi_zero_for_identical_distributions():
    props = [0.2] * 5
    assert psi(props, props) == pytest.approx(0.0, abs=1e-9)


def test_psi_positive_for_shifted_distribution():
    baseline = [0.2] * 5
    shifted = [0.05] * 4 + [0.80]
    assert psi(baseline, shifted) > 0.25


def test_exception_pack_writes_expected_rows():
    rules = load_rules()
    seeded = generate_seeded_snapshots()
    alerts = evaluate_stream(seeded, rules, population_size=COHORT_COUNT)

    with tempfile.TemporaryDirectory() as tmp:
        path = os.path.join(tmp, "exceptions.xlsx")
        written_path = write_exception_pack(alerts, path)
        assert written_path == path
        assert os.path.exists(path)

        wb = openpyxl.load_workbook(path)
        assert set(wb.sheetnames) == {
            "All exceptions", "cpr_tolerance", "stale_input", "population_stability",
        }

        cpr_alerts = [a for a in alerts if a["family"] == "cpr_tolerance"]
        assert cpr_alerts, "expected at least one cpr_tolerance alert to check against"
        first = cpr_alerts[0]

        ws = wb["cpr_tolerance"]
        header = [c.value for c in ws[1]]
        row = [c.value for c in ws[2]]
        row_dict = dict(zip(header, row))
        assert row_dict["rule_name"] == first["rule_name"]
        assert row_dict["cohort_id"] == first["cohort_id"]
        assert row_dict["predicted_cpr"] == first["snapshot"]["predicted_cpr"]
        assert row_dict["actual_cpr"] == first["snapshot"]["actual_cpr"]

        summary = wb["All exceptions"]
        assert summary["B2"].value == len(alerts)
